import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Document Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("PR Reviewer Portal & KT Chatbot")
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Backend REST API Specification & Workflow Architecture")
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run("Version 0.2.0 • FastAPI Backend • n8n Asynchronous Workflows")
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    doc.add_paragraph() # Spacer

    # Section 1: Overview
    h1 = doc.add_heading(level=1)
    run = h1.add_run("1. Executive Overview")
    run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_desc = doc.add_paragraph(
        "This specification documents the complete backend API suite for the PR Reviewer Portal. "
        "The system coordinates GitHub repository registration, automatic seeding of repository review rules "
        "and directory indexes, and an asynchronous AI-powered Knowledge Transfer (KT) Chatbot driven by n8n "
        "and Google Gemini 3.6 Flash."
    )
    p_desc.paragraph_format.line_spacing = 1.15

    # Overview Table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    meta_items = [
        ("Base URL (Host Port 1806)", "http://192.168.1.104:1806 (or http://localhost:1806)"),
        ("Base URL (Production)", "https://prreviewer.nik-server.in"),
        ("Interactive OpenAPI Docs", "/docs (Swagger UI) & /redoc"),
        ("Data Persistence", "SQLite (pr_reviewer.db) with Fernet symmetric credential encryption"),
        ("AI Orchestration Engine", "n8n (External Instance at https://n8n.nik-server.in)")
    ]

    for i, (k, v) in enumerate(meta_items):
        row = table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.6)
        c0.text = k
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(c0, "F7FAFC")
        set_cell_margins(c0, 60, 60, 100, 100)
        c1.text = v
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c1, 60, 60, 100, 100)

    doc.add_paragraph()

    # Section 2: Architecture & Security
    h1 = doc.add_heading(level=1)
    run = h1.add_run("2. Security & Credentials Architecture")
    run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    doc.add_paragraph(
        "To protect developer Personal Access Tokens (PAT) and App passwords while supporting background "
        "and asynchronous AI workflows, the backend implements dual-layer credential handling:"
    )

    bp1 = doc.add_paragraph(style='List Bullet')
    r1 = bp1.add_run("Integrity & Verification Layer: ")
    r1.bold = True
    bp1.add_run("The raw password is salt-hashed with scrypt (N=16384, r=8, p=1) and saved in repositories.app_password_hash. This one-way hash cannot be decrypted.")

    bp2 = doc.add_paragraph(style='List Bullet')
    r2 = bp2.add_run("Reversible Secure Storage Layer: ")
    r2.bold = True
    bp2.add_run("The password is also encrypted using AES-128-CBC via Fernet cryptography with the ENCRYPTION_SECRET_KEY and stored in repositories.app_password_encrypted. When the backend dispatches a query to n8n, it securely decrypts the PAT and transmits it over HTTPS.")

    doc.add_paragraph()

    # Section 3: API Endpoints
    h1 = doc.add_heading(level=1)
    run = h1.add_run("3. Complete API Endpoint Catalog")
    run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    apis = [
        {
            "tag": "System",
            "method": "GET",
            "path": "/health",
            "summary": "Health and Liveness Check",
            "desc": "Verifies that the FastAPI application and SQLite database connection are operational.",
            "request_body": "None",
            "params": "None",
            "responses": [
                ("200 OK", '{"status": "ok"}')
            ]
        },
        {
            "tag": "Repositories",
            "method": "POST",
            "path": "/api/repositories",
            "summary": "Verify and Register Repository",
            "desc": "Verifies access with GitHub, checks if the repository has already completed registration in green (returns 409 Conflict if all steps from file seeding to webhook registration are already active), seeds 'PR Reviewer/config.json', 'PR Reviewer/index.json', 'PR Reviewer/architecture-rules.md', and 'PR Reviewer/tech-stack.md' directly into the repository default branch, configures the GitHub pull_request webhook, and saves credentials in SQLite.",
            "request_body": '''{
  "repository_url": "https://github.com/org/repo",
  "github_username": "octocat",
  "github_app_password": "ghp_PersonalAccessTokenOrAppPassword"
}''',
            "params": "None",
            "responses": [
                ("201 Created", '''{
  "repository_url": "https://github.com/org/repo",
  "full_name": "org/repo",
  "registered_at": "2026-09-18T14:30:00.000000Z",
  "github_status": 200,
  "message": "Repository verified and registered."
}'''),
                ("400 Bad Request", '{"detail": "Invalid repository URL or GitHub rejected the request."}'),
                ("401 Unauthorized", '{"detail": "GitHub rejected the repository or token. Check both values and try again."}'),
                ("409 Conflict", '{"detail": "Repository \'org/repo\' is already registered. All setup steps from configuration to webhook registration have already been completed successfully."}'),
                ("502 Bad Gateway", '{"detail": "GitHub verification succeeded, but config.json could not be created."}')
            ]
        },
        {
            "tag": "Repositories",
            "method": "GET",
            "path": "/api/repositories",
            "summary": "List All Registered Repositories",
            "desc": "Returns all registered repositories. Safe for public/frontend display (no credentials or token hashes are returned). Includes registration status and webhook state.",
            "request_body": "None",
            "params": "None",
            "responses": [
                ("200 OK", '''[
  {
    "id": 1,
    "repository_url": "https://github.com/org/repo",
    "full_name": "org/repo",
    "github_username": "octocat",
    "status": "completed",
    "webhook_registered": true,
    "created_at": "2026-09-18 14:30:00"
  }
]''')
            ]
        },
        {
            "tag": "KT Chatbot",
            "method": "POST",
            "path": "/api/kt/chat",
            "summary": "Ask Question to Knowledge Transfer Chatbot",
            "desc": "Submits a question to the KT chatbot. Persists the question and an in-progress assistant record in chat_messages, decrypts the repository PAT, and dispatches an asynchronous execution request to n8n with the callback URL. Returns immediately to prevent HTTP timeouts.",
            "request_body": '''{
  "repository_id": 1,
  "session_id": "session-abc-123",
  "query": "Where is the authentication middleware and what architectural rules apply to it?"
}''',
            "params": "None",
            "responses": [
                ("202 Accepted", '''{
  "session_id": "session-abc-123",
  "status": "processing",
  "message": "Query submitted. The KT agent is generating the answer."
}'''),
                ("400 Bad Request", '{"detail": "Repository credentials missing or need re-registration."}'),
                ("404 Not Found", '{"detail": "Repository not found."}'),
                ("502 Bad Gateway", '{"detail": "Failed to communicate with n8n chatbot workflow: ..."}')
            ]
        },
        {
            "tag": "KT Chatbot",
            "method": "POST",
            "path": "/api/kt/chat/callback",
            "summary": "n8n Chat Completion Callback",
            "desc": "Invoked asynchronously by the n8n KT Chatbot workflow Callback node upon AI Agent completion. Updates the assistant record in SQLite from status='processing' to status='completed' with the generated answer.",
            "request_body": '''{
  "sessionId": "session-abc-123",
  "response": "# Authentication Architecture\\n\\nAuthentication is implemented in `app/auth/middleware.py`...\\n\\n```mermaid\\nflowchart TD\\n  Client --> Middleware\\n  Middleware --> TokenValidator\\n```",
  "status": "completed"
}''',
            "params": "None",
            "responses": [
                ("200 OK", '{"status": "ok"}')
            ]
        },
        {
            "tag": "KT Chatbot",
            "method": "GET",
            "path": "/api/kt/chat/messages",
            "summary": "Get Chat Message History",
            "desc": "Retrieves the complete chronologically ordered message thread for a conversation session. Used by the frontend for polling active queries and loading chat history.",
            "request_body": "None",
            "params": "session_id (Query String, Required, string)",
            "responses": [
                ("200 OK", '''[
  {
    "id": 1,
    "repository_id": 1,
    "session_id": "session-abc-123",
    "role": "user",
    "content": "Where is the authentication middleware?",
    "status": "completed",
    "created_at": "2026-09-18 14:32:00"
  },
  {
    "id": 2,
    "repository_id": 1,
    "session_id": "session-abc-123",
    "role": "assistant",
    "content": "# Authentication Architecture...",
    "status": "completed",
    "created_at": "2026-09-18 14:32:15"
  }
]''')
            ]
        },
        {
            "tag": "KT Chatbot",
            "method": "GET",
            "path": "/api/kt/chat/sessions",
            "summary": "List Chat Sessions by Repository",
            "desc": "Lists all conversation session threads initiated for a repository, including initial question preview and creation date.",
            "request_body": "None",
            "params": "repository_id (Query String, Required, integer)",
            "responses": [
                ("200 OK", '''[
  {
    "session_id": "session-abc-123",
    "first_query": "Where is the authentication middleware?",
    "created_at": "2026-09-18 14:32:00"
  }
]''')
            ]
        }
    ]

    for api in apis:
        h2 = doc.add_heading(level=2)
        r_tag = h2.add_run(f"[{api['tag']}] ")
        r_tag.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
        r_tag.font.size = Pt(11)

        r_method = h2.add_run(f"{api['method']} ")
        if api['method'] == 'GET':
            r_method.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0) # Blue
        elif api['method'] == 'POST':
            r_method.font.color.rgb = RGBColor(0x27, 0x67, 0x49) # Green

        r_path = h2.add_run(api['path'])
        r_path.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)

        p_sum = doc.add_paragraph()
        r_sum = p_sum.add_run(f"Summary: {api['summary']}")
        r_sum.font.bold = True
        p_sum.add_run(f"\n{api['desc']}")

        # Details Table
        tbl = doc.add_table(rows=0, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        def add_detail_row(label, content, is_code=False):
            row = tbl.add_row()
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Inches(1.8)
            c1.width = Inches(5.0)
            c0.text = label
            c0.paragraphs[0].runs[0].font.bold = True
            c0.paragraphs[0].runs[0].font.size = Pt(9)
            set_cell_background(c0, "EDF2F7")
            set_cell_margins(c0, 40, 40, 80, 80)
            c1.text = content
            r = c1.paragraphs[0].runs[0]
            r.font.size = Pt(9)
            if is_code:
                r.font.name = 'Consolas'
                set_cell_background(c1, "F7FAFC")
            else:
                set_cell_background(c1, "FFFFFF")
            set_cell_margins(c1, 40, 40, 80, 80)

        add_detail_row("Query Parameters", api['params'])
        if api['request_body'] != "None":
            add_detail_row("Request Body (JSON)", api['request_body'], is_code=True)
        for code, resp in api['responses']:
            add_detail_row(f"Response: {code}", resp, is_code=True)

        doc.add_paragraph() # Spacer

    # Section 4: Workflow Integration Architecture
    h1 = doc.add_heading(level=1)
    run = h1.add_run("4. n8n Workflow Integration Architecture")
    run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_wf = doc.add_paragraph(
        "Two dedicated n8n workflows complement the backend control plane:\n"
    )
    
    p_w1 = doc.add_paragraph(style='List Bullet')
    p_w1.add_run("Generate_KT_Docs.workflow.ts: ").bold = True
    p_w1.add_run("Batch workflow that analyzes the repository to generate initial markdown onboarding guides and Mermaid architecture flowcharts.")

    p_w2 = doc.add_paragraph(style='List Bullet')
    p_w2.add_run("KT_Chatbot.workflow.ts: ").bold = True
    p_w2.add_run("Conversational agent featuring:\n"
                 "• Webhook Trigger (POST /webhook/kt-chatbot, immediate 200 response)\n"
                 "• Google Gemini Chat Model (gemini-3.6-flash)\n"
                 "• Simple Memory (memoryBufferWindow bound to sessionId)\n"
                 "• GitHub API Tool (dynamic inspection of index.json, config.json, architecture rules, and codebase files)\n"
                 "• Callback Node (POST response back to /api/kt/chat/callback)")

    doc.add_paragraph()

    # Footer
    section = doc.sections[0]
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run("PR Reviewer Portal • Confidential & Proprietary Specification")
    f_run.font.size = Pt(8.5)
    f_run.font.color.rgb = RGBColor(0xA0, 0xAE, 0xC0)

    output_path = r"d:\hackathon project\portal\portal_apis_documentation.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()

